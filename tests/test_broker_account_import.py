from __future__ import annotations

import json
from pathlib import Path
import subprocess
import sys
import tempfile
import unittest

import pandas as pd
from openpyxl import Workbook

from scripts.data_pipeline.make_holdings_clean import (
    HABLE_ACTIVITY_COLUMNS,
    _canonical_ticker,
    _load_hable_account,
    broker_account_tool_recovery_metadata,
    import_broker_account_export,
)


class BrokerAccountImportTests(unittest.TestCase):
    def test_certified_alphanumeric_krx_security_id_is_preserved(self) -> None:
        self.assertEqual(_canonical_ticker("0009K0"), "0009K0")
        self.assertEqual(_canonical_ticker("A5930"), "005930")

    def add_hable_sheet(
        self,
        workbook: Workbook,
        *,
        title: str,
        account_id: str,
        account_asof: str = "2026-08-18",
        positions: list[dict[str, object]] | None = None,
        cash: int = 50_000,
        credit_or_loan: int = 0,
        withdrawable_cash: int | None = None,
        total_assets_override: int | None = None,
        holdings_summary_override: int | None = None,
    ) -> None:
        if title in workbook.sheetnames:
            sheet = workbook[title]
        else:
            sheet = workbook.create_sheet(title)
        positions = positions or [
            {
                "name": "삼성전자",
                "shares": 10,
                "average_price": 60_000,
                "market_price": 70_000,
                "commission": 100,
                "tax": 1_400,
            },
            {
                "name": "클래시스",
                "shares": 5,
                "average_price": 45_000,
                "market_price": 50_000,
                "commission": 50,
                "tax": 500,
            },
        ]

        sheet["A2"] = "실시간잔고(주식)"
        sheet["A4"] = "계좌번호:"
        sheet["B4"] = account_id
        sheet["S4"] = "출력일자:"
        year, month, day = account_asof.split("-")
        sheet["T4"] = f"{year}년{month}월{day}일"
        sheet["A5"] = "계좌명:"
        sheet["B5"] = "fixture account"
        sheet["S5"] = "출력시간:"
        sheet["T5"] = "09:30:00"
        sheet["A7"] = "추정자산 / D+2 추정예수금"
        sheet["E7"] = "총손익[①+②]"
        sheet["H7"] = "실현손익[①]"
        sheet["J7"] = "매입금액"
        sheet["K7"] = "평가금액"
        sheet["L7"] = "평가손익[②]"
        sheet["N7"] = "손익률"
        sheet["P7"] = "신용/대출금액"
        sheet["R7"] = "D+2 예수금"
        sheet["U7"] = "인출가능금액"
        sheet["A10"] = "종목명"
        sheet["D10"] = "평가손익"
        sheet["F10"] = "손익률"
        sheet["I10"] = "보유잔고"
        sheet["K10"] = "매도가능"
        sheet["M10"] = "평균단가"
        sheet["O10"] = "손익분기"
        sheet["Q10"] = "현재가"
        sheet["T10"] = "매입금액"
        sheet["A11"] = "평가금액"
        sheet["D11"] = "투자비중"
        sheet["F11"] = "수수료"
        sheet["I11"] = "제세금"
        sheet["K11"] = "대출일자"
        sheet["M11"] = "신용/대출금액"
        sheet["O11"] = "신용이자"
        sheet["Q11"] = "대비"
        sheet["T11"] = "등락률"

        holdings_value = 0
        purchase_value = 0
        total_commission = 0
        total_tax = 0
        for index, position in enumerate(positions):
            row = 12 + index * 2
            detail_row = row + 1
            shares = position["shares"]
            market_price = int(position["market_price"])
            average_price = int(position["average_price"])
            market_value = shares * market_price
            purchase_amount = shares * average_price
            commission = int(position["commission"])
            tax = int(position["tax"])
            sheet.cell(row, 1, position["name"])
            sheet.cell(row, 9, shares)
            sheet.cell(row, 11, shares)
            sheet.cell(row, 13, average_price)
            sheet.cell(row, 15, average_price)
            sheet.cell(row, 17, market_price)
            sheet.cell(row, 20, purchase_amount)
            sheet.cell(detail_row, 1, market_value)
            sheet.cell(detail_row, 6, commission)
            sheet.cell(detail_row, 9, tax)
            sheet.cell(detail_row, 17, "▼")
            holdings_value += market_value
            purchase_value += purchase_amount
            total_commission += commission
            total_tax += tax

        reported_holdings = holdings_value if holdings_summary_override is None else holdings_summary_override
        total_assets = (
            reported_holdings + cash - credit_or_loan - total_commission - total_tax
            if total_assets_override is None
            else total_assets_override
        )
        sheet["A8"] = total_assets
        sheet["C8"] = cash
        sheet["J8"] = purchase_value
        sheet["K8"] = reported_holdings
        sheet["P8"] = credit_or_loan
        sheet["R8"] = cash
        sheet["U8"] = max(cash, 0) if withdrawable_cash is None else withdrawable_cash

    def write_hable(
        self,
        root: Path,
        *,
        positions: list[dict[str, object]] | None = None,
        total_assets_override: int | None = None,
    ) -> Path:
        workbook = Workbook()
        self.add_hable_sheet(
            workbook,
            title="Sheet1",
            account_id="123-45-67890",
            positions=positions,
            total_assets_override=total_assets_override,
        )
        source = root / "explicit-user-selected-account.xlsx"
        workbook.save(source)
        return source

    @staticmethod
    def write_master(root: Path, rows: list[tuple[str, str]] | None = None) -> Path:
        master = root / "krx_master__fixture.csv"
        rows = rows or [("삼성전자", "005930"), ("클래시스", "214150")]
        pd.DataFrame(rows, columns=["name", "ticker"]).to_csv(
            master,
            index=False,
            encoding="utf-8-sig",
        )
        return master

    @staticmethod
    def write_reserve_classification(
        root: Path,
        *,
        account_asof: str = "2026-08-18",
        broker_name_alias: str = "RISE 미국단기투자등급회사채액?",
    ) -> Path:
        classification = root / "explicit-security-classification.json"
        classification.write_text(
            json.dumps(
                {
                    "contract_version": "ACCOUNT_ASSET_CLASSIFICATION_V1",
                    "account_asof": account_asof,
                    "classifications": [
                        {
                            "ticker": "437350",
                            "canonical_name": "RISE 미국단기투자등급회사채액티브 ETF",
                            "broker_name_alias": broker_name_alias,
                            "asset_type": "CASH_EQUIVALENT_RESERVE",
                            "reserve_flag": True,
                            "protected_keep_qty": True,
                            "counts_toward_cash_target": True,
                            "model_universe_eligible": False,
                            "holding_bonus_eligible": False,
                            "keep_current_eligible": False,
                            "top_k_eligible": False,
                            "equity_position_count_excluded": True,
                            "automatic_sell_prohibited_without_explicit_instruction": True,
                            "classification_source": "USER_EXPLICIT_CLASSIFICATION_FIXTURE",
                        }
                    ],
                },
                ensure_ascii=False,
                indent=2,
            ),
            encoding="utf-8",
        )
        return classification

    def import_fixture(
        self,
        root: Path,
        source: Path,
        *,
        master: Path | None = None,
        output_name: str = "account_import",
        account_id: str | None = None,
        account_asof: str = "2026-08-18",
        security_classification: Path | None = None,
    ) -> dict[str, object]:
        return import_broker_account_export(
            broker_account_file=source,
            output_dir=root / output_name,
            account_asof=account_asof,
            master_path=master or self.write_master(root),
            account_id=account_id,
            security_classification_path=security_classification,
        )

    def test_canonical_import_is_explicit_reconciled_and_private(self) -> None:
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            source = self.write_hable(root)
            result = self.import_fixture(root, source)
            holdings = pd.read_csv(result["holdings_path"], dtype={"ticker": str})
            self.assertEqual(holdings["ticker"].tolist(), ["005930", "214150"])
            self.assertTrue(holdings["ticker"].str.fullmatch(r"\d{6}").all())
            self.assertEqual(holdings["shares"].tolist(), [10, 5])
            snapshot = json.loads(Path(result["snapshot_path"]).read_text(encoding="utf-8"))
            self.assertEqual(snapshot["account_snapshot_status"], "VERIFIED")
            self.assertEqual(snapshot["cash_balance_signed"], 50_000)
            self.assertIsNone(snapshot["orderable_cash"])
            self.assertEqual(snapshot["orderable_cash_status"], "NOT_AVAILABLE")
            self.assertEqual(snapshot["gross_account_nav"], 1_000_000)
            self.assertEqual(snapshot["liquidation_nav"], 997_950)
            self.assertEqual(snapshot["broker_total_assets"], snapshot["liquidation_nav"])
            self.assertEqual(snapshot["reconciliation"]["error_krw"], 0)
            self.assertEqual(
                snapshot["liquidation_reconciliation"]["status"],
                "PRESERVED_BROKER_ESTIMATE_NOT_USED_FOR_TARGET_WEIGHTS",
            )
            self.assertFalse(snapshot["reconciliation"]["synthetic_cash_created"])

            activity = pd.read_csv(result["activity_path"], dtype=str).fillna("")
            self.assertEqual(activity.columns.tolist(), HABLE_ACTIVITY_COLUMNS)
            self.assertEqual(activity.loc[0, "event_type"], "NOT_AVAILABLE")
            manifest_text = Path(result["manifest_path"]).read_text(encoding="utf-8")
            manifest = json.loads(manifest_text)
            self.assertNotIn("1234567890", manifest_text)
            self.assertNotIn(source.name, manifest_text)
            self.assertNotIn(str(source.resolve()), manifest_text)
            self.assertEqual(manifest["execution_input_source"], "BROKER_ACCOUNT_EXPORT_EXPLICIT")
            self.assertFalse(manifest["execution_plan_eligible"])
            self.assertEqual(
                manifest["execution_plan_blocking_failures"],
                ["ORDERABLE_CASH_NOT_AVAILABLE"],
            )
            self.assertFalse(manifest["synthetic_cash_used"])
            for artifact in manifest["canonical_artifacts"].values():
                self.assertEqual(Path(artifact["path"]).name, artifact["path"])
                self.assertFalse(Path(artifact["path"]).is_absolute())

    def test_20260820_layout_preserves_reserve_metadata_and_separates_nav_bases(self) -> None:
        positions = [
            {"name": "삼성전자", "shares": 100, "average_price": 53_800, "market_price": 273_000, "commission": 10_540, "tax": 54_200},
            {"name": "메디아나", "shares": 151, "average_price": 19_127, "market_price": 9_220, "commission": 630, "tax": 2_784},
            {"name": "카페24", "shares": 106, "average_price": 27_000, "market_price": 16_850, "commission": 680, "tax": 3_582},
            {"name": "큐렉소", "shares": 174, "average_price": 16_330, "market_price": 9_010, "commission": 650, "tax": 3_135},
            {"name": "티에스이", "shares": 27, "average_price": 105_900, "market_price": 242_500, "commission": 1_400, "tax": 13_068},
            {"name": "선익시스템", "shares": 33, "average_price": 86_500, "market_price": 67_000, "commission": 750, "tax": 4_428},
            {"name": "클래시스", "shares": 600, "average_price": 23_270, "market_price": 33_600, "commission": 10_100, "tax": 40_320},
            {"name": "RFHIC", "shares": 37, "average_price": 77_900, "market_price": 50_100, "commission": 700, "tax": 3_707},
            {"name": "티에프이", "shares": 46, "average_price": 61_900, "market_price": 40_800, "commission": 700, "tax": 3_735},
            {"name": "RISE 미국단기투자등급회사채액?", "shares": 1_050, "average_price": 10_346, "market_price": 12_335, "commission": 4_716, "tax": 0},
        ]
        master_rows = [
            ("삼성전자", "005930"),
            ("메디아나", "041920"),
            ("카페24", "042000"),
            ("큐렉소", "060280"),
            ("티에스이", "131290"),
            ("선익시스템", "171090"),
            ("클래시스", "214150"),
            ("RFHIC", "218410"),
            ("티에프이", "425420"),
        ]
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            workbook = Workbook()
            self.add_hable_sheet(
                workbook,
                title="Sheet1",
                account_id="123-45-67890",
                account_asof="2026-08-20",
                positions=positions,
                cash=1_466_753,
                credit_or_loan=0,
                withdrawable_cash=0,
                total_assets_override=78_972_575,
            )
            source = root / "explicit-current-account.xlsx"
            workbook.save(source)
            classification = self.write_reserve_classification(root, account_asof="2026-08-20")
            result = self.import_fixture(
                root,
                source,
                master=self.write_master(root, master_rows),
                account_asof="2026-08-20",
                security_classification=classification,
            )

            holdings = pd.read_csv(result["holdings_path"], dtype={"ticker": str})
            self.assertEqual(len(holdings), 10)
            reserve = holdings.loc[holdings["ticker"] == "437350"].iloc[0]
            self.assertEqual(reserve["name"], "RISE 미국단기투자등급회사채액티브 ETF")
            self.assertEqual(int(reserve["shares"]), 1_050)
            self.assertEqual(reserve["asset_type"], "CASH_EQUIVALENT_RESERVE")
            for field in [
                "reserve_flag",
                "protected_keep_qty",
                "counts_toward_cash_target",
                "equity_position_count_excluded",
                "automatic_sell_prohibited_without_explicit_instruction",
            ]:
                self.assertTrue(bool(reserve[field]))
            for field in [
                "model_universe_eligible",
                "holding_bonus_eligible",
                "keep_current_eligible",
                "top_k_eligible",
            ]:
                self.assertFalse(bool(reserve[field]))

            snapshot = result["snapshot"]
            self.assertEqual(snapshot["broker_holdings_market_value"], 77_646_810)
            self.assertEqual(snapshot["cash_balance_signed"], 1_466_753)
            self.assertEqual(snapshot["credit_or_loan_amount"], 0)
            self.assertEqual(snapshot["gross_account_nav"], 79_113_563)
            self.assertEqual(snapshot["liquidation_nav"], 78_972_575)
            self.assertEqual(snapshot["broker_total_assets"], 78_972_575)
            self.assertIsNone(snapshot["orderable_cash"])
            self.assertEqual(snapshot["broker_withdrawable_cash"], 0)
            self.assertEqual(snapshot["reconciliation"]["status"], "PASS")
            self.assertEqual(snapshot["reconciliation"]["error_krw"], 0)
            self.assertFalse(snapshot["liquidation_reconciliation"]["planning_nav_eligible"])
            self.assertEqual(
                snapshot["liquidation_reconciliation"]["implied_broker_liquidation_adjustment_krw"],
                -140_988,
            )
            security_audit = result["manifest"]["security_classification"]
            self.assertTrue(security_audit["provided"])
            self.assertEqual(security_audit["matched_tickers"], ["437350"])
            self.assertFalse(security_audit["source_path_stored"])

    def test_explicit_reserve_alias_and_classification_asof_are_fail_closed(self) -> None:
        position = {
            "name": "RISE 미국단기투자등급회사채액?",
            "shares": 1_050,
            "average_price": 10_346,
            "market_price": 12_335,
            "commission": 4_716,
            "tax": 0,
        }
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            source = self.write_hable(root, positions=[position])
            wrong_alias = self.write_reserve_classification(
                root,
                broker_name_alias="RISE 미국단기투자등급회사채액티브",
            )
            with self.assertRaisesRegex(ValueError, "exact alias"):
                self.import_fixture(
                    root,
                    source,
                    master=self.write_master(root, [("삼성전자", "005930")]),
                    security_classification=wrong_alias,
                )
            self.assertFalse((root / "account_import").exists())

        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            source = self.write_hable(root, positions=[position])
            wrong_asof = self.write_reserve_classification(root, account_asof="2026-08-19")
            with self.assertRaisesRegex(ValueError, "account_asof"):
                self.import_fixture(
                    root,
                    source,
                    master=self.write_master(root, [("삼성전자", "005930")]),
                    security_classification=wrong_asof,
                )
            self.assertFalse((root / "account_import").exists())

    def test_account_asof_is_explicit_and_not_inferred_from_filename(self) -> None:
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            source = self.write_hable(root)
            with self.assertRaisesRegex(ValueError, "내장 출력일자"):
                import_broker_account_export(
                    broker_account_file=source,
                    output_dir=root / "out",
                    account_asof="1999-01-01",
                    master_path=self.write_master(root),
                )
            self.assertFalse((root / "out").exists())

    def test_production_cli_never_auto_selects_a_broker_file(self) -> None:
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            completed = subprocess.run(
                [
                    sys.executable,
                    "scripts/data_pipeline/make_holdings_clean.py",
                    "--output-dir",
                    str(root / "account_import"),
                    "--account-asof",
                    "2026-08-18",
                    "--master",
                    str(self.write_master(root)),
                ],
                cwd=Path(__file__).resolve().parents[1],
                text=True,
                capture_output=True,
                check=False,
            )
            self.assertNotEqual(completed.returncode, 0)
            self.assertIn("--broker-account-file", completed.stderr)
            self.assertFalse((root / "account_import").exists())

    def test_production_cli_exact_arguments_create_canonical_artifacts(self) -> None:
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            source = self.write_hable(root)
            output_dir = root / "account_import"
            completed = subprocess.run(
                [
                    sys.executable,
                    "scripts/data_pipeline/make_holdings_clean.py",
                    "--broker-account-file",
                    str(source),
                    "--output-dir",
                    str(output_dir),
                    "--account-asof",
                    "2026-08-18",
                    "--account-id",
                    "1234567890",
                    "--master",
                    str(self.write_master(root)),
                ],
                cwd=Path(__file__).resolve().parents[1],
                text=True,
                capture_output=True,
                check=False,
            )
            self.assertEqual(completed.returncode, 0, completed.stderr)
            self.assertTrue((output_dir / "holdings_clean__account_asof=2026-08-18.csv").exists())
            self.assertTrue((output_dir / "account_snapshot__account_asof=2026-08-18.json").exists())
            self.assertTrue((output_dir / "account_activity__from=NOT_AVAILABLE__to=2026-08-18.csv").exists())
            self.assertTrue((output_dir / "broker_account_import_manifest.json").exists())

    def test_multiple_accounts_require_selector_and_selector_is_masked(self) -> None:
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            workbook = Workbook()
            self.add_hable_sheet(workbook, title="Sheet1", account_id="111-11-11111")
            self.add_hable_sheet(workbook, title="SecondAccount", account_id="222-22-22222")
            source = root / "multi.xlsx"
            workbook.save(source)
            master = self.write_master(root)
            with self.assertRaisesRegex(ValueError, "복수 계좌"):
                self.import_fixture(root, source, master=master)
            result = self.import_fixture(
                root,
                source,
                master=master,
                output_name="selected",
                account_id="2222222222",
            )
            manifest_text = Path(result["manifest_path"]).read_text(encoding="utf-8")
            self.assertNotIn("2222222222", manifest_text)
            self.assertEqual(result["manifest"]["account_count_detected"], 2)
            self.assertTrue(result["manifest"]["account_selector_provided"])

    def test_duplicate_ticker_is_summed_with_audit(self) -> None:
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            positions = [
                {
                    "name": "삼성전자",
                    "shares": 2,
                    "average_price": 60_000,
                    "market_price": 70_000,
                    "commission": 20,
                    "tax": 280,
                },
                {
                    "name": "삼성전자",
                    "shares": 3,
                    "average_price": 65_000,
                    "market_price": 70_000,
                    "commission": 30,
                    "tax": 420,
                },
            ]
            source = self.write_hable(root, positions=positions)
            result = self.import_fixture(root, source, master=self.write_master(root, [("삼성전자", "005930")]))
            holdings = pd.read_csv(result["holdings_path"], dtype={"ticker": str})
            self.assertEqual(holdings.loc[0, "ticker"], "005930")
            self.assertEqual(int(holdings.loc[0, "shares"]), 5)
            audit = result["manifest"]["duplicate_ticker_aggregation"]
            self.assertEqual(audit[0]["source_row_count"], 2)
            self.assertEqual(audit[0]["summed_shares"], 5)

    def test_negative_and_fractional_shares_fail_closed(self) -> None:
        base = {
            "name": "삼성전자",
            "average_price": 60_000,
            "market_price": 70_000,
            "commission": 0,
            "tax": 0,
        }
        for shares, message in [(-1, "음수"), (1.5, "소수")]:
            with self.subTest(shares=shares), tempfile.TemporaryDirectory() as tmp:
                root = Path(tmp)
                source = self.write_hable(root, positions=[{**base, "shares": shares}])
                with self.assertRaisesRegex(ValueError, message):
                    self.import_fixture(root, source, master=self.write_master(root, [("삼성전자", "005930")]))

    def test_unsupported_product_and_holdings_reconciliation_mismatch_fail_closed(self) -> None:
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            unsupported = {
                "name": "지원하지않는ETF",
                "shares": 1,
                "average_price": 10_000,
                "market_price": 10_000,
                "commission": 0,
                "tax": 0,
            }
            source = self.write_hable(root, positions=[unsupported])
            with self.assertRaisesRegex(ValueError, "지원하지 않는 상품"):
                self.import_fixture(root, source, master=self.write_master(root))
            self.assertFalse((root / "account_import").exists())

        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            workbook = Workbook()
            self.add_hable_sheet(
                workbook,
                title="Sheet1",
                account_id="123-45-67890",
                holdings_summary_override=1,
            )
            source = root / "summary-mismatch.xlsx"
            workbook.save(source)
            with self.assertRaisesRegex(ValueError, "요약 보유평가액"):
                self.import_fixture(root, source)
            self.assertFalse((root / "account_import").exists())

    def test_unknown_nonempty_sheet_fails_closed(self) -> None:
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            workbook = Workbook()
            self.add_hable_sheet(workbook, title="Sheet1", account_id="123-45-67890")
            notes = workbook.create_sheet("Unknown")
            notes["A1"] = "unrecognized layout"
            source = root / "unknown-sheet.xlsx"
            workbook.save(source)
            with self.assertRaisesRegex(ValueError, "인증되지 않은 non-empty sheet"):
                self.import_fixture(root, source)

    def test_legacy_input_output_cli_is_preserved(self) -> None:
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            source = root / "legacy.csv"
            pd.DataFrame(
                [
                    ["종목명", "보유수량"],
                    ["", ""],
                    ["삼성전자", "3"],
                ]
            ).to_csv(source, index=False, header=False, encoding="utf-8-sig")
            master = self.write_master(root, [("삼성전자", "005930")])
            output = root / "clean.csv"
            completed = subprocess.run(
                [
                    sys.executable,
                    "scripts/data_pipeline/make_holdings_clean.py",
                    "--input",
                    str(source),
                    "--output",
                    str(output),
                    "--master",
                    str(master),
                ],
                cwd=Path(__file__).resolve().parents[1],
                text=True,
                capture_output=True,
                check=False,
            )
            self.assertEqual(completed.returncode, 0, completed.stderr)
            clean = pd.read_csv(output, dtype={"ticker": str})
            self.assertEqual(clean.to_dict("records"), [{"ticker": "005930", "name": "삼성전자", "shares": 3}])

    def test_recovery_metadata_and_historical_layout_provenance(self) -> None:
        metadata = broker_account_tool_recovery_metadata()
        self.assertEqual(
            metadata["original_introduced_commit"],
            "230579ae4a8db5ab85e38705361d65ed2a9e6f24",
        )
        self.assertEqual(
            metadata["restored_commit"],
            "f9f654a1df9f9ef9b1ebb6a6539a1dc2b34e361d",
        )
        historical = Path("data/portfolio/history/20260330.xlsx")
        if historical.exists():
            _, sheet, masked, embedded, count = _load_hable_account(
                historical,
                account_id=None,
                sheet_name=None,
            )
            self.assertEqual(sheet, "Sheet1")
            self.assertEqual(embedded.isoformat(), "2026-03-30")
            self.assertEqual(count, 1)
            self.assertIn("*", masked)


if __name__ == "__main__":
    unittest.main()
