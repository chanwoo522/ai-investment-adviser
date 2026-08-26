from __future__ import annotations

import argparse
from datetime import date, datetime, timezone
from decimal import Decimal, InvalidOperation
import hashlib
import json
from pathlib import Path
import re
import sys
from typing import Any, Iterable

import pandas as pd

try:
    from scripts.common.security_id import InvalidSecurityId, normalize_security_id
except ModuleNotFoundError:  # Preserve direct ``python scripts/...`` CLI usage.
    sys.path.insert(0, str(Path(__file__).resolve().parents[2]))
    from scripts.common.security_id import InvalidSecurityId, normalize_security_id


NAME_CANDIDATES = ["종목명", "종목 명", "한글명", "주식명", "종목"]
SHARES_CANDIDATES = ["보유잔고", "보유수량", "수량", "잔고수량", "잔고"]

# 마스터 파일 컬럼 후보
MASTER_NAME_CANDIDATES = ["종목명", "한글종목명", "주식종목명", "name", "name_final"]
MASTER_TICKER_CANDIDATES = ["종목코드", "ticker", "단축코드", "티커", "code"]

PARSER_NAME = "scripts.data_pipeline.make_holdings_clean"
PARSER_VERSION = "2.1.0-production-canonical"
HABLE_BROKER_NAME = "H-able"
SECURITY_CLASSIFICATION_CONTRACT = "ACCOUNT_ASSET_CLASSIFICATION_V1"
SECURITY_CLASSIFICATION_BOOL_FIELDS = [
    "reserve_flag",
    "protected_keep_qty",
    "counts_toward_cash_target",
    "model_universe_eligible",
    "holding_bonus_eligible",
    "keep_current_eligible",
    "top_k_eligible",
    "equity_position_count_excluded",
    "automatic_sell_prohibited_without_explicit_instruction",
]
SECURITY_METADATA_FIELDS = ["asset_type", *SECURITY_CLASSIFICATION_BOOL_FIELDS]
HABLE_ACTIVITY_COLUMNS = [
    "event_datetime",
    "settlement_date",
    "event_type",
    "ticker",
    "name",
    "quantity",
    "price",
    "gross_amount",
    "commission",
    "transaction_tax",
    "other_fee",
    "net_cashflow",
    "external_cashflow_flag",
    "source_row_id",
]

_RECOVERY_METADATA: dict[str, Any] = {
    "status": "RECOVERED_EXISTING_TOOL_EXTENDED",
    "original_path": "scripts/_deprecated_20260330/make_holdings_clean.py",
    "original_introduced_commit": "230579ae4a8db5ab85e38705361d65ed2a9e6f24",
    "original_blob": "dbeaa7e0704697b7c590b8a2176a84d32c63482d",
    "mainline_deleted_commit": "3f78cfcd8e819570b6503a56a6cd0d730ea53862",
    "restored_path": "scripts/data_pipeline/make_holdings_clean.py",
    "restored_commit": "f9f654a1df9f9ef9b1ebb6a6539a1dc2b34e361d",
    "restored_blob": "02214d26378c973743797a43c8f51c04fb37c66f",
    "legacy_cli_preserved": [
        "--input",
        "--output",
        "--master",
        "--asof",
        "--sheet_name",
        "--master_sheet_name",
        "--allow_missing_ticker",
    ],
    "production_extension": "HABLE_EXACT_SNAPSHOT_LAYOUT_V1",
    "explicit_security_classification_contract": SECURITY_CLASSIFICATION_CONTRACT,
}


def broker_account_tool_recovery_metadata() -> dict[str, Any]:
    """Return non-sensitive Git provenance for the recovered account tool."""

    return json.loads(json.dumps(_RECOVERY_METADATA))


def _norm_text(x: object) -> str:
    if pd.isna(x):
        return ""
    s = str(x)
    s = s.replace("\u3000", " ")
    s = re.sub(r"\s+", "", s)
    return s.strip()


def _normalize_name_key(x: object) -> str:
    if pd.isna(x):
        return ""
    s = str(x).strip()
    s = s.replace("\u3000", " ")
    s = re.sub(r"\s+", "", s)
    return s.upper()


def _normalize_ticker_value(x: object) -> str | None:
    if pd.isna(x):
        return None
    s = re.sub(r"\D", "", str(x))
    if not s:
        return None
    return s.zfill(6)


def _find_first_matching_col(columns: Iterable[object], candidates: list[str]) -> str | None:
    normalized = {str(c): _norm_text(c) for c in columns}
    cand_norms = [_norm_text(c) for c in candidates]
    for orig, norm in normalized.items():
        if any(c == norm or c in norm for c in cand_norms):
            return orig
    return None


def _detect_header_start(raw: pd.DataFrame) -> int:
    max_rows = min(len(raw), 60)
    name_cands = [_norm_text(x) for x in NAME_CANDIDATES]
    shares_cands = [_norm_text(x) for x in SHARES_CANDIDATES]

    for i in range(max_rows):
        row = [_norm_text(v) for v in raw.iloc[i].tolist()]
        rowset = set(v for v in row if v)

        has_name = any(any(c in v for c in name_cands) for v in rowset)
        has_shares = any(any(c in v for c in shares_cands) for v in rowset)

        if has_name and has_shares:
            return i

    raise ValueError("보유내역 헤더 시작점을 찾지 못했습니다. 원본 파일 형식을 확인해 주세요.")


def _build_headers(h1: list[object], h2: list[object]) -> list[str]:
    headers: list[str] = []
    for a, b in zip(h1, h2):
        a2 = _norm_text(a)
        b2 = _norm_text(b)
        if a2 and b2:
            headers.append(f"{a2}_{b2}")
        elif a2:
            headers.append(a2)
        elif b2:
            headers.append(b2)
        else:
            headers.append("")
    return headers


def _load_raw_table(input_path: Path, sheet_name: str | int | None = None) -> pd.DataFrame:
    suffix = input_path.suffix.lower()
    if suffix in {".xls", ".xlsx", ".xlsm"}:
        kwargs = {"header": None}
        if sheet_name is not None:
            kwargs["sheet_name"] = sheet_name
        return pd.read_excel(input_path, **kwargs)
    return pd.read_csv(input_path, header=None, encoding="utf-8-sig")


def _coerce_int_series(s: pd.Series) -> pd.Series:
    s = (
        s.astype(str)
        .str.replace(",", "", regex=False)
        .str.replace(" ", "", regex=False)
        .str.replace("\u3000", "", regex=False)
    )
    s = s.replace({"": pd.NA, "nan": pd.NA, "None": pd.NA, "<NA>": pd.NA})
    return pd.to_numeric(s, errors="coerce").round().astype("Int64")


def _extract_holdings(raw: pd.DataFrame) -> pd.DataFrame:
    start_idx = _detect_header_start(raw)

    if start_idx + 1 >= len(raw):
        raise ValueError("헤더 2행 구조를 읽을 수 없습니다.")

    header1 = raw.iloc[start_idx].tolist()
    header2 = raw.iloc[start_idx + 1].tolist()
    headers = _build_headers(header1, header2)

    data = raw.iloc[start_idx + 2 :].copy().reset_index(drop=True)
    data.columns = headers

    name_col = _find_first_matching_col(data.columns, NAME_CANDIDATES)
    shares_col = _find_first_matching_col(data.columns, SHARES_CANDIDATES)

    if name_col is None:
        raise ValueError(f"종목명 컬럼을 찾지 못했습니다. columns={list(data.columns)}")
    if shares_col is None:
        raise ValueError(f"보유수량 컬럼을 찾지 못했습니다. columns={list(data.columns)}")

    extracted = data[[name_col, shares_col]].copy()
    extracted.columns = ["name", "shares"]

    extracted["name"] = extracted["name"].map(lambda x: str(x).strip() if not pd.isna(x) else "")
    extracted["shares"] = _coerce_int_series(extracted["shares"])

    # 빈값 제거
    extracted = extracted[(extracted["name"] != "") & extracted["shares"].notna()].copy()

    # 숫자만 있는 행 제거
    extracted = extracted[
        ~extracted["name"].str.fullmatch(r"\d+(?:\.\d+)?", na=False)
    ].copy()

    # 보유수량 0 초과만 유지
    extracted = extracted[extracted["shares"] > 0].copy()

    # 동일 종목이 여러 번 있으면 합산
    extracted = extracted.groupby("name", as_index=False)["shares"].sum()

    return extracted


def _load_master(master_path: Path, sheet_name: str | int | None = None) -> pd.DataFrame:
    if not master_path.exists():
        raise FileNotFoundError(f"master not found: {master_path}")

    suffix = master_path.suffix.lower()
    if suffix == ".parquet":
        m = pd.read_parquet(master_path)
    elif suffix in {".xls", ".xlsx", ".xlsm"}:
        kwargs = {}
        if sheet_name is not None:
            kwargs["sheet_name"] = sheet_name
        m = pd.read_excel(master_path, **kwargs)
    else:
        m = pd.read_csv(master_path, encoding="utf-8-sig")

    print(f"[INFO] master columns: {list(m.columns)}")

    name_col = _find_first_matching_col(m.columns, MASTER_NAME_CANDIDATES)
    ticker_col = _find_first_matching_col(m.columns, MASTER_TICKER_CANDIDATES)

    print(f"[INFO] selected master name col: {name_col}")
    print(f"[INFO] selected master ticker col: {ticker_col}")

    if name_col is None:
        raise ValueError(f"master에서 종목명 컬럼을 찾지 못했습니다. columns={list(m.columns)}")
    if ticker_col is None:
        raise ValueError(f"master에서 ticker 컬럼을 찾지 못했습니다. columns={list(m.columns)}")

    out = m[[name_col, ticker_col]].copy()
    out.columns = ["name", "ticker"]

    out["name"] = out["name"].astype(str).str.strip()
    out["ticker"] = out["ticker"].map(_normalize_ticker_value)
    out["name_key"] = out["name"].map(_normalize_name_key)

    out = out.dropna(subset=["ticker"])
    out = out[out["name"] != ""].copy()
    out = out[out["name_key"] != ""].copy()
    out = out.drop_duplicates(subset=["name_key"], keep="first").reset_index(drop=True)

    return out


def _auto_find_master(asof: str | None) -> Path:
    # 기본 fallback. 가능하면 --master를 직접 주는 것을 권장.
    candidates: list[Path] = []
    p1 = Path("reference")
    p2 = Path("data/processed")

    if asof:
        candidates.extend(sorted(p1.glob(f"*{asof.replace('-', '')}*.xlsx"), reverse=True))
        candidates.extend(sorted(p1.glob(f"*{asof}*.xlsx"), reverse=True))
        candidates.extend(sorted(p2.glob(f"krx_master__asof={asof}__src=*__v=*.parquet"), reverse=True))
        candidates.extend(sorted(p2.glob(f"krx_master__asof={asof}__src=*__v=*.csv"), reverse=True))

    candidates.extend(sorted(p1.glob("*.xlsx"), reverse=True))
    candidates.extend(sorted(p2.glob("krx_master__asof=*__src=*__v=*.parquet"), reverse=True))
    candidates.extend(sorted(p2.glob("krx_master__asof=*__src=*__v=*.csv"), reverse=True))

    if not candidates:
        raise FileNotFoundError("reference 또는 data/processed 아래에서 master 파일을 찾지 못했습니다.")
    return candidates[0]


def _sha256_file(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def _is_missing(value: object) -> bool:
    if value is None:
        return True
    try:
        return bool(pd.isna(value))
    except (TypeError, ValueError):
        return False


def _decimal_value(value: object, *, label: str) -> Decimal:
    if _is_missing(value):
        raise ValueError(f"H-able 필수 금액이 없습니다: {label}")
    if isinstance(value, bool):
        raise ValueError(f"H-able 금액 형식이 올바르지 않습니다: {label}")
    text = str(value).strip().replace(",", "")
    try:
        parsed = Decimal(text)
    except InvalidOperation as exc:
        raise ValueError(f"H-able 금액 형식이 올바르지 않습니다: {label}") from exc
    if not parsed.is_finite():
        raise ValueError(f"H-able 금액이 유한수가 아닙니다: {label}")
    return parsed


def _krw_integer(value: object, *, label: str) -> int:
    parsed = _decimal_value(value, label=label)
    if parsed != parsed.to_integral_value():
        raise ValueError(f"원 단위 정수가 아닌 H-able 금액입니다: {label}")
    return int(parsed)


def _strict_share_integer(value: object, *, source_row_id: str) -> int:
    parsed = _decimal_value(value, label=f"shares ({source_row_id})")
    if parsed != parsed.to_integral_value():
        raise ValueError(f"소수 주식은 production import에서 허용되지 않습니다: {source_row_id}")
    shares = int(parsed)
    if shares < 0:
        raise ValueError(f"음수 보유수량은 production import에서 허용되지 않습니다: {source_row_id}")
    return shares


def _canonical_ticker(value: object) -> str:
    if _is_missing(value):
        raise ValueError("master ticker가 비어 있습니다.")
    text = str(value).strip().replace(",", "").upper()
    # Preserve the importer's legacy acceptance of an exchange-style ``A``
    # prefix for numeric codes, while allowing certified six-character KRX
    # identifiers such as ``0009K0`` without extracting or destroying letters.
    if re.fullmatch(r"A\d{1,6}", text):
        text = text[1:]
    try:
        return str(normalize_security_id(text))
    except InvalidSecurityId as exc:
        raise ValueError(
            "master ticker는 1~6자리 숫자 또는 인증된 6자리 영숫자 KRX ID여야 합니다."
        ) from exc


def _normalize_account_id(value: object) -> str:
    if _is_missing(value):
        raise ValueError("H-able 계좌번호가 비어 있습니다.")
    normalized = re.sub(r"[^0-9A-Za-z]", "", str(value)).upper()
    if not normalized:
        raise ValueError("H-able 계좌번호를 정규화할 수 없습니다.")
    return normalized


def _mask_account_id(value: object) -> str:
    normalized = _normalize_account_id(value)
    if len(normalized) <= 4:
        return "*" * len(normalized)
    return "*" * (len(normalized) - 4) + normalized[-4:]


def _label_key(value: object) -> str:
    return _norm_text(value).rstrip(":：")


def _table_value(table: pd.DataFrame, row: int, col: int) -> object:
    row_idx = row - 1
    col_idx = col - 1
    if row_idx < 0 or col_idx < 0 or row_idx >= len(table) or col_idx >= len(table.columns):
        return None
    return table.iat[row_idx, col_idx]


def _table_is_empty(table: pd.DataFrame) -> bool:
    if table.empty:
        return True
    return all(_is_missing(value) for value in table.to_numpy().ravel().tolist())


def _hable_signature_issues(table: pd.DataFrame) -> list[str]:
    expected = {
        (2, 1): "실시간잔고(주식)",
        (4, 1): "계좌번호",
        (5, 1): "계좌명",
        (7, 1): "추정자산/D+2추정예수금",
        (7, 11): "평가금액",
        (7, 18): "D+2예수금",
        (7, 21): "인출가능금액",
        (10, 1): "종목명",
        (10, 9): "보유잔고",
        (11, 6): "수수료",
        (11, 9): "제세금",
    }
    issues: list[str] = []
    for (row, col), label in expected.items():
        actual = _label_key(_table_value(table, row, col))
        if actual != _label_key(label):
            issues.append(f"R{row}C{col}:{label}")
    return issues


def _parse_hable_output_date(value: object) -> date:
    match = re.fullmatch(r"\s*(\d{4})년\s*(\d{1,2})월\s*(\d{1,2})일\s*", str(value))
    if match is None:
        raise ValueError("H-able 출력일자 형식이 인증된 구조와 다릅니다.")
    return date(int(match.group(1)), int(match.group(2)), int(match.group(3)))


def _worksheet_table(worksheet: Any) -> pd.DataFrame:
    rows = [list(row) for row in worksheet.iter_rows(values_only=True)]
    return pd.DataFrame(rows)


def _load_hable_account(
    source_path: Path,
    *,
    account_id: str | None,
    sheet_name: str | int | None,
) -> tuple[pd.DataFrame, str, str, date, int]:
    if source_path.suffix.lower() != ".xlsx":
        raise ValueError("production account snapshot은 인증된 H-able .xlsx 형식만 지원합니다.")

    try:
        from openpyxl import load_workbook
    except ImportError as exc:
        raise RuntimeError("H-able xlsx import에는 openpyxl이 필요합니다.") from exc

    try:
        workbook = load_workbook(source_path, read_only=True, data_only=True)
    except Exception as exc:
        raise ValueError("H-able xlsx를 읽을 수 없습니다.") from exc

    candidates: list[dict[str, Any]] = []
    unknown_nonempty_sheets: list[str] = []
    workbook_sheet_names = list(workbook.sheetnames)
    for index, worksheet in enumerate(workbook.worksheets):
        table = _worksheet_table(worksheet)
        if _table_is_empty(table):
            continue
        issues = _hable_signature_issues(table)
        if issues:
            unknown_nonempty_sheets.append(worksheet.title)
            continue
        account_labels = sum(
            1 for value in table.to_numpy().ravel().tolist() if _label_key(value) == "계좌번호"
        )
        if account_labels != 1:
            raise ValueError("한 sheet에서 복수 계좌가 탐지되어 인증된 H-able 구조로 처리할 수 없습니다.")
        raw_account_id = _table_value(table, 4, 2)
        candidates.append(
            {
                "table": table,
                "sheet_name": worksheet.title,
                "sheet_index": index,
                "raw_account_id": raw_account_id,
                "normalized_account_id": _normalize_account_id(raw_account_id),
                "output_date": _parse_hable_output_date(_table_value(table, 4, 20)),
            }
        )

    workbook.close()
    if unknown_nonempty_sheets:
        joined = ", ".join(sorted(unknown_nonempty_sheets))
        raise ValueError(f"인증되지 않은 non-empty sheet가 있습니다: {joined}")
    if not candidates:
        raise ValueError("인증된 H-able account snapshot sheet를 찾지 못했습니다.")
    normalized_ids = [str(candidate["normalized_account_id"]) for candidate in candidates]
    if len(set(normalized_ids)) != len(normalized_ids):
        raise ValueError("동일 계좌가 여러 sheet에서 중복 탐지되었습니다.")
    if len(candidates) > 1 and not account_id:
        raise ValueError("복수 계좌가 탐지되었습니다. --account-id를 명시해야 합니다.")

    if account_id:
        selector = _normalize_account_id(account_id)
        selected = [candidate for candidate in candidates if candidate["normalized_account_id"] == selector]
        if len(selected) != 1:
            raise ValueError("--account-id가 탐지된 계좌와 일치하지 않습니다.")
        candidate = selected[0]
    else:
        candidate = candidates[0]

    if sheet_name is not None:
        if isinstance(sheet_name, int):
            try:
                constrained_name = workbook_sheet_names[sheet_name]
            except IndexError as exc:
                raise ValueError("--sheet_name index가 workbook 범위를 벗어났습니다.") from exc
        else:
            constrained_name = str(sheet_name)
        if candidate["sheet_name"] != constrained_name:
            raise ValueError("--sheet_name이 선택된 --account-id sheet와 일치하지 않습니다.")

    return (
        candidate["table"],
        str(candidate["sheet_name"]),
        _mask_account_id(candidate["raw_account_id"]),
        candidate["output_date"],
        len(candidates),
    )


def _load_master_table_strict(
    master_path: Path,
    *,
    sheet_name: str | int | None,
) -> pd.DataFrame:
    if not master_path.exists():
        raise FileNotFoundError(f"master not found: {master_path}")
    suffix = master_path.suffix.lower()
    if suffix == ".parquet":
        master = pd.read_parquet(master_path)
    elif suffix in {".xls", ".xlsx", ".xlsm"}:
        kwargs: dict[str, object] = {}
        if sheet_name is not None:
            kwargs["sheet_name"] = sheet_name
        master = pd.read_excel(master_path, **kwargs)
    elif suffix == ".csv":
        master = pd.read_csv(master_path, encoding="utf-8-sig", dtype="object")
    else:
        raise ValueError("production master 형식은 parquet/xls/xlsx/xlsm/csv만 지원합니다.")

    name_col = _find_first_matching_col(master.columns, MASTER_NAME_CANDIDATES)
    ticker_col = _find_first_matching_col(master.columns, MASTER_TICKER_CANDIDATES)
    if name_col is None or ticker_col is None:
        raise ValueError("master에 인증 가능한 종목명/ticker 컬럼이 없습니다.")
    mapped = master[[name_col, ticker_col]].copy()
    mapped.columns = ["name", "ticker"]
    mapped["name"] = mapped["name"].map(lambda value: "" if _is_missing(value) else str(value).strip())
    mapped = mapped[mapped["name"] != ""].copy()
    try:
        mapped["ticker"] = mapped["ticker"].map(_canonical_ticker)
    except ValueError as exc:
        raise ValueError("master에 인증 불가능한 ticker가 있습니다.") from exc
    mapped["name_key"] = mapped["name"].map(_normalize_name_key)
    mapped = mapped[mapped["name_key"] != ""].copy()
    ambiguity = mapped.groupby("name_key")["ticker"].nunique()
    if bool((ambiguity > 1).any()):
        raise ValueError("동일 종목명이 master의 복수 ticker에 매핑됩니다.")
    return mapped.drop_duplicates(["name_key", "ticker"], keep="first").reset_index(drop=True)


def _load_security_classifications(
    classification_path: Path | None,
    *,
    account_asof: date,
) -> tuple[pd.DataFrame, dict[str, Any]]:
    columns = [
        "ticker",
        "canonical_name",
        "broker_name_alias",
        *SECURITY_METADATA_FIELDS,
    ]
    if classification_path is None:
        return pd.DataFrame(columns=columns), {
            "provided": False,
            "contract_version": None,
            "source_sha256": None,
            "classification_count": 0,
            "source_path_stored": False,
        }

    path = Path(classification_path)
    if not path.exists() or not path.is_file():
        raise FileNotFoundError(f"security classification not found: {path}")
    if path.suffix.lower() != ".json":
        raise ValueError("--security-classification은 인증된 JSON 형식만 지원합니다.")
    try:
        payload = json.loads(path.read_text(encoding="utf-8"))
    except (OSError, UnicodeDecodeError, json.JSONDecodeError) as exc:
        raise ValueError("security classification JSON을 읽을 수 없습니다.") from exc
    if not isinstance(payload, dict):
        raise ValueError("security classification 최상위 값은 object여야 합니다.")
    if payload.get("contract_version") != SECURITY_CLASSIFICATION_CONTRACT:
        raise ValueError("security classification contract_version이 지원되지 않습니다.")
    if payload.get("account_asof") != account_asof.isoformat():
        raise ValueError("security classification account_asof가 명시적 계좌 기준일과 다릅니다.")
    raw_rows = payload.get("classifications")
    if not isinstance(raw_rows, list) or not raw_rows:
        raise ValueError("security classification classifications가 비어 있습니다.")

    required = set(columns)
    allowed = required | {"classification_source"}
    rows: list[dict[str, Any]] = []
    for index, raw in enumerate(raw_rows):
        if not isinstance(raw, dict):
            raise ValueError(f"security classification row가 object가 아닙니다: {index}")
        missing = sorted(required - set(raw))
        unknown = sorted(set(raw) - allowed)
        if missing or unknown:
            raise ValueError(
                f"security classification schema 불일치 row={index} missing={missing} unknown={unknown}"
            )
        ticker = _canonical_ticker(raw["ticker"])
        canonical_name = str(raw["canonical_name"])
        broker_name_alias = str(raw["broker_name_alias"])
        if not canonical_name or canonical_name != canonical_name.strip():
            raise ValueError("canonical_name은 앞뒤 공백 없는 비어 있지 않은 문자열이어야 합니다.")
        if not broker_name_alias or broker_name_alias != broker_name_alias.strip():
            raise ValueError("broker_name_alias는 앞뒤 공백 없는 정확한 문자열이어야 합니다.")
        if raw["asset_type"] != "CASH_EQUIVALENT_RESERVE":
            raise ValueError("현재 explicit classification은 CASH_EQUIVALENT_RESERVE만 인증합니다.")
        for field in SECURITY_CLASSIFICATION_BOOL_FIELDS:
            if type(raw[field]) is not bool:
                raise ValueError(f"security classification {field}는 boolean이어야 합니다.")
        expected_reserve_contract = {
            "reserve_flag": True,
            "protected_keep_qty": True,
            "counts_toward_cash_target": True,
            "model_universe_eligible": False,
            "holding_bonus_eligible": False,
            "keep_current_eligible": False,
            "top_k_eligible": False,
            "equity_position_count_excluded": True,
            "automatic_sell_prohibited_without_explicit_instruction": True,
        }
        if any(raw[field] is not expected for field, expected in expected_reserve_contract.items()):
            raise ValueError("CASH_EQUIVALENT_RESERVE eligibility/protection 계약이 올바르지 않습니다.")
        row = {
            "ticker": ticker,
            "canonical_name": canonical_name,
            "broker_name_alias": broker_name_alias,
            "classification_source": str(
                raw.get("classification_source", "EXPLICIT_SECURITY_CLASSIFICATION")
            ),
        }
        row.update({field: raw[field] for field in SECURITY_METADATA_FIELDS})
        rows.append(row)

    frame = pd.DataFrame(rows)
    if bool(frame["broker_name_alias"].duplicated().any()):
        raise ValueError("security classification broker_name_alias가 중복됩니다.")
    if bool(frame["ticker"].duplicated().any()):
        raise ValueError("security classification ticker가 중복됩니다.")
    return frame, {
        "provided": True,
        "contract_version": SECURITY_CLASSIFICATION_CONTRACT,
        "source_sha256": _sha256_file(path),
        "classification_count": int(len(frame)),
        "source_path_stored": False,
    }


def _excel_column_name(index_zero_based: int) -> str:
    number = index_zero_based + 1
    result = ""
    while number:
        number, remainder = divmod(number - 1, 26)
        result = chr(65 + remainder) + result
    return result


def _extract_hable_positions(table: pd.DataFrame, *, sheet_name: str) -> tuple[list[dict[str, Any]], dict[str, Any]]:
    header_idx = _detect_header_start(table)
    if header_idx + 1 >= len(table):
        raise ValueError("H-able 2행 보유내역 헤더가 완전하지 않습니다.")
    headers = _build_headers(table.iloc[header_idx].tolist(), table.iloc[header_idx + 1].tolist())
    name_col = _find_first_matching_col(headers, NAME_CANDIDATES)
    shares_col = _find_first_matching_col(headers, SHARES_CANDIDATES)
    if name_col is None or shares_col is None:
        raise ValueError("H-able 종목명/보유수량 컬럼을 찾지 못했습니다.")
    name_idx = headers.index(name_col)
    shares_idx = headers.index(shares_col)
    excluded_rows: list[dict[str, Any]] = []
    positions: list[dict[str, Any]] = []

    for row_idx in range(header_idx + 2, len(table)):
        name_value = table.iat[row_idx, name_idx]
        if _is_missing(name_value):
            continue
        name_text = str(name_value).strip()
        if not name_text or re.fullmatch(r"[-+]?\d+(?:\.\d+)?", name_text.replace(",", "")):
            continue
        source_row_id = f"{sheet_name}!{row_idx + 1}"
        shares = _strict_share_integer(table.iat[row_idx, shares_idx], source_row_id=source_row_id)
        if shares == 0:
            excluded_rows.append({"source_row_id": source_row_id, "reason": "ZERO_SHARES"})
            continue
        detail_idx = row_idx + 1
        if detail_idx >= len(table):
            raise ValueError(f"H-able 종목 상세행이 없습니다: {source_row_id}")
        detail_name = table.iat[detail_idx, name_idx]
        if _is_missing(detail_name) or re.fullmatch(r"[-+]?\d+(?:\.\d+)?", str(detail_name).replace(",", "")) is None:
            raise ValueError(f"H-able 종목 상세행 구조가 인증본과 다릅니다: {source_row_id}")

        average_price = _krw_integer(_table_value(table, row_idx + 1, 13), label=f"average price {source_row_id}")
        market_price = _krw_integer(_table_value(table, row_idx + 1, 17), label=f"market price {source_row_id}")
        market_value = _krw_integer(_table_value(table, detail_idx + 1, 1), label=f"market value {source_row_id}")
        commission = _krw_integer(_table_value(table, detail_idx + 1, 6), label=f"commission {source_row_id}")
        transaction_tax = _krw_integer(_table_value(table, detail_idx + 1, 9), label=f"tax {source_row_id}")
        if min(average_price, market_price, market_value, commission, transaction_tax) < 0:
            raise ValueError(f"H-able 종목 금액/비용에 음수가 있습니다: {source_row_id}")
        if shares * market_price != market_value:
            raise ValueError(f"H-able 종목 평가금액이 수량×현재가와 일치하지 않습니다: {source_row_id}")
        positions.append(
            {
                "name": name_text,
                "name_key": _normalize_name_key(name_text),
                "shares": shares,
                "broker_market_price": market_price,
                "broker_market_value": market_value,
                "average_purchase_price": average_price,
                "estimated_liquidation_commission": commission,
                "estimated_liquidation_tax": transaction_tax,
                "source_row_id": source_row_id,
            }
        )

    if not positions:
        raise ValueError("H-able account snapshot에 양수 보유종목이 없습니다.")
    mapping = {
        "sheet": sheet_name,
        "header_row_one_based": header_idx + 1,
        "name": {"source_header": name_col, "column": _excel_column_name(name_idx)},
        "shares": {"source_header": shares_col, "column": _excel_column_name(shares_idx)},
        "broker_market_price": {"source_header": "현재가", "column": "Q"},
        "broker_market_value": {"source_header": "평가금액", "column": "A", "row": "detail"},
        "average_purchase_price": {"source_header": "평균단가", "column": "M"},
        "estimated_liquidation_commission": {"source_header": "수수료", "column": "F", "row": "detail"},
        "estimated_liquidation_tax": {"source_header": "제세금", "column": "I", "row": "detail"},
    }
    return positions, {"excluded_rows": excluded_rows, "column_mapping": mapping}


def _map_and_aggregate_positions(
    positions: list[dict[str, Any]],
    master: pd.DataFrame,
    *,
    account_id_masked: str,
    security_classifications: pd.DataFrame | None = None,
) -> tuple[pd.DataFrame, pd.DataFrame, list[dict[str, Any]], list[dict[str, Any]]]:
    raw = pd.DataFrame(positions)
    mapped = raw.merge(master[["name_key", "ticker"]], on="name_key", how="left", validate="many_to_one")
    mapped["canonical_name"] = mapped["name"]
    mapped["security_classification_source"] = "KRX_EQUITY_MASTER"
    equity_metadata: dict[str, Any] = {
        "asset_type": "EQUITY",
        "reserve_flag": False,
        "protected_keep_qty": False,
        "counts_toward_cash_target": False,
        "model_universe_eligible": True,
        "holding_bonus_eligible": True,
        "keep_current_eligible": True,
        "top_k_eligible": True,
        "equity_position_count_excluded": False,
        "automatic_sell_prohibited_without_explicit_instruction": False,
    }
    for field, value in equity_metadata.items():
        mapped[field] = value

    classification_matches: list[dict[str, Any]] = []
    classifications = security_classifications
    if classifications is not None and not classifications.empty:
        alias_to_index = {
            str(row.broker_name_alias): index
            for index, row in classifications.iterrows()
        }
        used_classification_indexes: set[int] = set()
        for row_index, position in mapped.iterrows():
            broker_name = str(position["name"])
            classification_index = alias_to_index.get(broker_name)
            if classification_index is None:
                continue
            classification = classifications.loc[classification_index]
            classified_ticker = _canonical_ticker(classification["ticker"])
            existing_ticker = position["ticker"]
            if not _is_missing(existing_ticker) and _canonical_ticker(existing_ticker) != classified_ticker:
                raise ValueError("equity master와 explicit security classification ticker가 충돌합니다.")
            mapped.at[row_index, "ticker"] = classified_ticker
            mapped.at[row_index, "canonical_name"] = str(classification["canonical_name"])
            mapped.at[row_index, "security_classification_source"] = str(
                classification["classification_source"]
            )
            for field in SECURITY_METADATA_FIELDS:
                mapped.at[row_index, field] = classification[field]
            used_classification_indexes.add(int(classification_index))
            classification_matches.append(
                {
                    "ticker": classified_ticker,
                    "broker_name_alias": broker_name,
                    "source_row_id": str(position["source_row_id"]),
                    "match_basis": "EXACT_BROKER_NAME_ALIAS",
                }
            )
        if len(used_classification_indexes) != len(classifications):
            raise ValueError("명시적 security classification 중 계좌파일과 exact alias 매칭되지 않은 행이 있습니다.")

    if bool(mapped["ticker"].isna().any()):
        raise ValueError("지원하지 않는 상품 또는 ticker mapping 누락 종목이 있습니다.")
    mapped["ticker"] = mapped["ticker"].map(_canonical_ticker)

    output_rows: list[dict[str, Any]] = []
    duplicate_audit: list[dict[str, Any]] = []
    for ticker, group in mapped.groupby("ticker", sort=True):
        canonical_names = set(group["canonical_name"].astype(str).tolist())
        if len(canonical_names) != 1:
            raise ValueError("동일 ticker에 서로 다른 canonical 종목명이 매핑되었습니다.")
        prices = set(int(value) for value in group["broker_market_price"].tolist())
        if len(prices) != 1:
            raise ValueError("동일 ticker 중복행의 broker market price가 다릅니다.")
        metadata_values = {
            field: set(group[field].tolist())
            for field in SECURITY_METADATA_FIELDS
        }
        if any(len(values) != 1 for values in metadata_values.values()):
            raise ValueError("동일 ticker 중복행의 asset metadata가 다릅니다.")
        shares = int(group["shares"].sum())
        weighted_average = sum(
            int(row.average_purchase_price) * int(row.shares) for row in group.itertuples(index=False)
        ) / shares
        output_rows.append(
            {
                "ticker": _canonical_ticker(ticker),
                "name": str(group.iloc[0]["canonical_name"]),
                "shares": shares,
                "broker_market_price": int(group.iloc[0]["broker_market_price"]),
                "broker_market_value": int(group["broker_market_value"].sum()),
                "average_purchase_price": weighted_average,
                "account_id_masked": account_id_masked,
                **{field: group.iloc[0][field] for field in SECURITY_METADATA_FIELDS},
                "security_classification_source": str(
                    group.iloc[0]["security_classification_source"]
                ),
            }
        )
        if len(group) > 1:
            duplicate_audit.append(
                {
                    "ticker": _canonical_ticker(ticker),
                    "source_row_count": int(len(group)),
                    "summed_shares": shares,
                    "source_row_ids": sorted(group["source_row_id"].astype(str).tolist()),
                    "aggregation_basis": "SAME_TICKER_SAME_NORMALIZED_NAME",
                }
            )
    holdings = pd.DataFrame(output_rows).sort_values("ticker").reset_index(drop=True)
    if not bool(holdings["ticker"].astype(str).str.fullmatch(r"\d{6}").all()):
        raise ValueError("canonical holdings ticker가 6자리 문자열이 아닙니다.")
    if not bool(holdings["shares"].map(lambda value: isinstance(value, int) and value > 0).all()):
        raise ValueError("canonical holdings shares가 양의 정수가 아닙니다.")
    return holdings, mapped, duplicate_audit, classification_matches


def _build_hable_snapshot(
    table: pd.DataFrame,
    mapped_positions: pd.DataFrame,
    *,
    account_asof: date,
    account_id_masked: str,
    source_sha256: str,
    imported_at: str,
) -> dict[str, Any]:
    broker_total_assets = _krw_integer(_table_value(table, 8, 1), label="broker total assets A8")
    broker_d2_cash_alt = _krw_integer(_table_value(table, 8, 3), label="D+2 estimated cash C8")
    broker_holdings_market_value = _krw_integer(_table_value(table, 8, 11), label="holdings market value K8")
    credit_loan_amount = _krw_integer(_table_value(table, 8, 16), label="credit/loan amount P8")
    cash_balance_signed = _krw_integer(_table_value(table, 8, 18), label="D+2 cash R8")
    withdrawable_cash = _krw_integer(_table_value(table, 8, 21), label="withdrawable cash U8")
    if broker_d2_cash_alt != cash_balance_signed:
        raise ValueError("H-able의 두 D+2 예수금 표시값이 일치하지 않습니다.")

    detail_holdings_market_value = int(mapped_positions["broker_market_value"].sum())
    if abs(detail_holdings_market_value - broker_holdings_market_value) > 1:
        raise ValueError("H-able 요약 보유평가액과 종목별 평가액 합계가 일치하지 않습니다.")
    liquidation_commission = int(mapped_positions["estimated_liquidation_commission"].sum())
    liquidation_tax = int(mapped_positions["estimated_liquidation_tax"].sum())
    gross_account_nav = (
        broker_holdings_market_value
        + cash_balance_signed
        - credit_loan_amount
    )
    calculated_liquidation_nav = (
        gross_account_nav
        - liquidation_commission
        - liquidation_tax
    )
    liquidation_reconciliation_error = calculated_liquidation_nav - broker_total_assets

    return {
        "account_asof": account_asof.isoformat(),
        "account_id_masked": account_id_masked,
        "currency": "KRW",
        "broker_name": HABLE_BROKER_NAME,
        "broker_total_assets": broker_total_assets,
        "broker_total_assets_basis": "BROKER_ESTIMATED_ASSETS_LIQUIDATION_NAV",
        "gross_account_nav": gross_account_nav,
        "gross_account_nav_basis": "HOLDINGS_MARKET_VALUE_PLUS_SIGNED_CASH_MINUS_LOAN",
        "liquidation_nav": broker_total_assets,
        "liquidation_nav_basis": "BROKER_ESTIMATED_ASSETS_AFTER_DISPLAYED_LIQUIDATION_COSTS",
        "broker_holdings_market_value": broker_holdings_market_value,
        "cash_balance_signed": cash_balance_signed,
        "orderable_cash": None,
        "orderable_cash_status": "NOT_AVAILABLE",
        "settlement_receivable": None,
        "settlement_payable": None,
        "broker_withdrawable_cash": withdrawable_cash,
        "credit_or_loan_amount": credit_loan_amount,
        "estimated_liquidation_commission": liquidation_commission,
        "estimated_liquidation_tax": liquidation_tax,
        "explicit_other_accounts": [
            {"name": "CREDIT_OR_LOAN_LIABILITY", "amount": -credit_loan_amount},
            {"name": "DISPLAYED_LIQUIDATION_COMMISSION", "amount": -liquidation_commission},
            {"name": "DISPLAYED_LIQUIDATION_TAX", "amount": -liquidation_tax},
        ],
        "source_file_sha256": source_sha256,
        "parser_name": PARSER_NAME,
        "parser_version": PARSER_VERSION,
        "imported_at": imported_at,
        "account_snapshot_status": "VERIFIED",
        "execution_input_source": "BROKER_ACCOUNT_EXPORT_EXPLICIT",
        "reconciliation": {
            "formula": "broker_holdings_market_value + cash_balance_signed - credit_or_loan_amount",
            "detail_holdings_market_value": detail_holdings_market_value,
            "calculated_gross_account_nav": gross_account_nav,
            "gross_account_nav": gross_account_nav,
            "error_krw": detail_holdings_market_value - broker_holdings_market_value,
            "tolerance_krw": 1,
            "status": "PASS",
            "synthetic_cash_created": False,
        },
        "liquidation_reconciliation": {
            "comparison_formula": "gross_account_nav - displayed_liquidation_commission - displayed_liquidation_tax",
            "comparison_liquidation_nav": calculated_liquidation_nav,
            "reported_liquidation_nav": broker_total_assets,
            "implied_broker_liquidation_adjustment_krw": broker_total_assets - gross_account_nav,
            "difference_vs_displayed_cost_components_krw": -liquidation_reconciliation_error,
            "status": "PRESERVED_BROKER_ESTIMATE_NOT_USED_FOR_TARGET_WEIGHTS",
            "planning_nav_eligible": False,
        },
    }


def import_broker_account_export(
    *,
    broker_account_file: Path,
    output_dir: Path,
    account_asof: str,
    master_path: Path,
    account_id: str | None = None,
    sheet_name: str | int | None = None,
    master_sheet_name: str | int | None = None,
    security_classification_path: Path | None = None,
) -> dict[str, Any]:
    """Import one explicitly selected H-able account export into canonical artifacts.

    This production path never searches for an account file and never derives the
    account date from its filename.  The explicit date must match the workbook's
    embedded output date.
    """

    source_path = Path(broker_account_file)
    destination = Path(output_dir)
    if not source_path.exists() or not source_path.is_file():
        raise FileNotFoundError(f"broker account file not found: {source_path}")
    if destination.exists():
        raise FileExistsError("immutable account import output directory already exists")
    try:
        explicit_asof = date.fromisoformat(str(account_asof))
    except ValueError as exc:
        raise ValueError("--account-asof는 YYYY-MM-DD 형식이어야 합니다.") from exc

    table, selected_sheet, masked_account, embedded_date, account_count = _load_hable_account(
        source_path,
        account_id=account_id,
        sheet_name=sheet_name,
    )
    if embedded_date != explicit_asof:
        raise ValueError("--account-asof가 H-able 내장 출력일자와 일치하지 않습니다.")
    master = _load_master_table_strict(Path(master_path), sheet_name=master_sheet_name)
    security_classifications, classification_provenance = _load_security_classifications(
        Path(security_classification_path) if security_classification_path is not None else None,
        account_asof=explicit_asof,
    )
    positions, extraction_audit = _extract_hable_positions(table, sheet_name=selected_sheet)
    holdings, mapped_positions, duplicate_audit, classification_matches = _map_and_aggregate_positions(
        positions,
        master,
        account_id_masked=masked_account,
        security_classifications=security_classifications,
    )

    imported_at = datetime.now(timezone.utc).isoformat()
    source_sha256 = _sha256_file(source_path)
    snapshot = _build_hable_snapshot(
        table,
        mapped_positions,
        account_asof=explicit_asof,
        account_id_masked=masked_account,
        source_sha256=source_sha256,
        imported_at=imported_at,
    )
    activity = pd.DataFrame([{column: "" for column in HABLE_ACTIVITY_COLUMNS}])
    activity.loc[0, "event_type"] = "NOT_AVAILABLE"
    activity.loc[0, "source_row_id"] = "NOT_AVAILABLE"

    holdings_name = f"holdings_clean__account_asof={explicit_asof.isoformat()}.csv"
    snapshot_name = f"account_snapshot__account_asof={explicit_asof.isoformat()}.json"
    activity_name = f"account_activity__from=NOT_AVAILABLE__to={explicit_asof.isoformat()}.csv"
    manifest_name = "broker_account_import_manifest.json"

    destination.mkdir(parents=True, exist_ok=False)
    holdings_path = destination / holdings_name
    snapshot_path = destination / snapshot_name
    activity_path = destination / activity_name
    manifest_path = destination / manifest_name
    holdings.to_csv(holdings_path, index=False, encoding="utf-8-sig")
    snapshot_path.write_text(json.dumps(snapshot, ensure_ascii=False, indent=2), encoding="utf-8")
    activity.to_csv(activity_path, index=False, encoding="utf-8-sig")

    canonical_artifacts = {
        "holdings": {"path": holdings_name, "sha256": _sha256_file(holdings_path)},
        "account_snapshot": {"path": snapshot_name, "sha256": _sha256_file(snapshot_path)},
        "account_activity": {"path": activity_name, "sha256": _sha256_file(activity_path)},
    }
    manifest: dict[str, Any] = {
        "contract_version": "BROKER_ACCOUNT_IMPORT_V1",
        "status": "PASS",
        "production_eligible_account_snapshot": True,
        "account_snapshot_status": "VERIFIED",
        "execution_plan_eligible": False,
        "execution_plan_blocking_failures": ["ORDERABLE_CASH_NOT_AVAILABLE"],
        "execution_input_source": "BROKER_ACCOUNT_EXPORT_EXPLICIT",
        "source_file_sha256": source_sha256,
        "source_file_format": source_path.suffix.lower().lstrip("."),
        "source_path_stored": False,
        "account_asof": explicit_asof.isoformat(),
        "account_id_masked": masked_account,
        "account_count_detected": account_count,
        "account_selector_provided": account_id is not None,
        "activity_ledger_status": "NOT_AVAILABLE",
        "activity_start_date": None,
        "activity_end_date": None,
        "sheet_and_column_mapping": extraction_audit["column_mapping"],
        "security_classification": {
            **classification_provenance,
            "matched_tickers": sorted(
                {str(match["ticker"]) for match in classification_matches}
            ),
            "matches": classification_matches,
        },
        "duplicate_ticker_aggregation": duplicate_audit,
        "excluded_rows": extraction_audit["excluded_rows"],
        "warnings": [
            "ACCOUNT_ACTIVITY_NOT_AVAILABLE",
            "ORDERABLE_CASH_NOT_AVAILABLE_WITHDRAWABLE_CASH_NOT_SUBSTITUTED",
        ],
        "canonical_artifacts": canonical_artifacts,
        "parser_provenance": {
            "parser_name": PARSER_NAME,
            "parser_version": PARSER_VERSION,
            "layout_contract": "HABLE_EXACT_SNAPSHOT_LAYOUT_V1",
            "tool_recovery": broker_account_tool_recovery_metadata(),
        },
        "master_sha256": _sha256_file(Path(master_path)),
        "master_path_stored": False,
        "snapshot_reconciliation": snapshot["reconciliation"],
        "liquidation_reconciliation": snapshot["liquidation_reconciliation"],
        "synthetic_cash_used": False,
        "legacy_available_cash_used": False,
        "legacy_execution_nav_used": False,
        "conditional_performance_nav_used": False,
        "imported_at": imported_at,
    }
    manifest_path.write_text(json.dumps(manifest, ensure_ascii=False, indent=2), encoding="utf-8")
    return {
        "holdings_path": holdings_path,
        "snapshot_path": snapshot_path,
        "activity_path": activity_path,
        "manifest_path": manifest_path,
        "snapshot": snapshot,
        "manifest": manifest,
    }


def _sheet_selector(value: str | None) -> str | int | None:
    if value is None:
        return None
    try:
        return int(value)
    except ValueError:
        return value


def _run_legacy_holdings_clean(args: argparse.Namespace) -> None:
    if not args.input or not args.output:
        raise ValueError("legacy mode에는 --input과 --output이 모두 필요합니다.")
    if args.output_dir or args.account_asof or args.account_id:
        raise ValueError("legacy --input mode와 production account 옵션을 혼합할 수 없습니다.")

    input_path = Path(args.input)
    output_path = Path(args.output)
    if not input_path.exists():
        raise FileNotFoundError(f"input holdings file not found: {input_path}")
    sheet_name = _sheet_selector(args.sheet_name)
    master_sheet_name = _sheet_selector(args.master_sheet_name)
    master_path = Path(args.master) if args.master else _auto_find_master(args.asof)

    print(f"[INFO] input : {input_path}")
    print(f"[INFO] master: {master_path}")
    raw = _load_raw_table(input_path, sheet_name=sheet_name)
    holdings = _extract_holdings(raw)
    master = _load_master(master_path, sheet_name=master_sheet_name)
    holdings["name_key"] = holdings["name"].map(_normalize_name_key)
    out = holdings.merge(master[["name_key", "ticker"]], on="name_key", how="left").drop(columns=["name_key"])

    missing = out["ticker"].isna().sum()
    print(f"[INFO] parsed holdings rows: {len(out)}")
    print(f"[INFO] missing ticker rows : {int(missing)}")
    if missing > 0:
        print("[WARN] ticker 매핑 실패 종목:")
        print(out.loc[out["ticker"].isna(), ["name", "shares"]].to_string(index=False))
        if not args.allow_missing_ticker:
            raise ValueError("ticker 매핑 실패 종목이 있습니다. --allow_missing_ticker 없이 계속할 수 없습니다.")

    out = out[["ticker", "name", "shares"]].copy()
    out["shares"] = pd.to_numeric(out["shares"], errors="coerce").round().astype("Int64")
    out = out.dropna(subset=["shares"])
    out = out.sort_values(["ticker", "name"], na_position="last").reset_index(drop=True)
    output_path.parent.mkdir(parents=True, exist_ok=True)
    out.to_csv(output_path, index=False, encoding="utf-8-sig")
    print(f"[OK] saved: {output_path}")
    print("\n[PREVIEW]")
    print(out.head(30).to_string(index=False))


def main() -> None:
    ap = argparse.ArgumentParser(description="증권사 원본 holdings 파일(xls/xlsx/csv)을 actions용 clean csv로 변환")
    ap.add_argument("--input", default=None, help="legacy 원본 holdings 파일 경로 (.xls/.xlsx/.csv)")
    ap.add_argument("--output", default=None, help="legacy clean csv 출력 경로")
    ap.add_argument("--master", default=None, help="종목명-티커 매핑 master 파일 경로 (.xlsx/.csv/.parquet)")
    ap.add_argument("--asof", default=None, help="master 자동 탐색 시 우선 사용할 asof")
    ap.add_argument("--sheet_name", default=None, help="input 엑셀 시트명 또는 시트 index")
    ap.add_argument("--master_sheet_name", default=None, help="master 엑셀 시트명 또는 시트 index")
    ap.add_argument("--allow_missing_ticker", action="store_true", help="ticker 누락 종목이 있어도 계속 진행")
    ap.add_argument("--broker-account-file", default=None, help="명시적으로 선택한 증권사 계좌 다운로드 파일")
    ap.add_argument("--output-dir", default=None, help="새 immutable canonical account import 디렉터리")
    ap.add_argument("--account-asof", default=None, help="계좌파일 기준일 YYYY-MM-DD (파일명 추론 금지)")
    ap.add_argument("--account-id", default=None, help="복수 계좌 선택자 (로그/산출물에는 마스킹)")
    ap.add_argument(
        "--security-classification",
        default=None,
        help="master 밖 상품을 exact broker alias로 매핑하는 명시적 JSON",
    )
    args = ap.parse_args()

    if args.broker_account_file:
        if args.input or args.output:
            raise ValueError("--broker-account-file과 legacy --input/--output을 혼합할 수 없습니다.")
        if args.allow_missing_ticker:
            raise ValueError("production account import에서는 --allow_missing_ticker를 사용할 수 없습니다.")
        if not args.output_dir or not args.account_asof:
            raise ValueError("production account import에는 --output-dir과 --account-asof가 필요합니다.")
        if not args.master:
            raise ValueError("production account import에는 명시적인 --master가 필요합니다.")
        result = import_broker_account_export(
            broker_account_file=Path(args.broker_account_file),
            output_dir=Path(args.output_dir),
            account_asof=args.account_asof,
            master_path=Path(args.master),
            account_id=args.account_id,
            sheet_name=_sheet_selector(args.sheet_name),
            master_sheet_name=_sheet_selector(args.master_sheet_name),
            security_classification_path=(
                Path(args.security_classification) if args.security_classification else None
            ),
        )
        snapshot = result["snapshot"]
        print(f"[OK] account snapshot status: {snapshot['account_snapshot_status']}")
        print(f"[OK] account id: {snapshot['account_id_masked']}")
        print(f"[OK] account asof: {snapshot['account_asof']}")
        print(f"[OUT] {Path(result['holdings_path']).name}")
        print(f"[OUT] {Path(result['snapshot_path']).name}")
        print(f"[OUT] {Path(result['activity_path']).name}")
        print(f"[OUT] {Path(result['manifest_path']).name}")
        return

    if args.output_dir or args.account_asof or args.account_id or args.security_classification:
        raise ValueError("--broker-account-file 없이 production account 옵션을 사용할 수 없습니다.")
    _run_legacy_holdings_clean(args)


if __name__ == "__main__":
    try:
        main()
    except Exception as e:
        print(f"[ERROR] {e}", file=sys.stderr)
        raise
